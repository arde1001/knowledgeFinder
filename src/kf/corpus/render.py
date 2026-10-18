"""Write one FileSpec plus its prose paragraphs to disk in the file's format."""

from pathlib import Path

_ASCII = {"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-", "…": "...",
          "€": "EUR", " ": " ", "•": "-"}


def _latin1(text):
    for k, v in _ASCII.items():
        text = text.replace(k, v)
    return text.encode("latin-1", "replace").decode("latin-1")


def render(spec, paragraphs, root):
    path = Path(root) / spec.path
    path.parent.mkdir(parents=True, exist_ok=True)
    WRITERS[spec.fmt](spec, paragraphs, path)


def _md(spec, paragraphs, path):
    parts = [f"# {spec.topic}"]
    if spec.table:
        head, *rows = spec.table
        parts.append("\n".join(["| " + " | ".join(map(str, head)) + " |", "|" + "---|" * len(head)]
                               + ["| " + " | ".join(map(str, r)) + " |" for r in rows]))
    path.write_text("\n\n".join(parts + paragraphs) + "\n", encoding="utf-8")


def _txt(spec, paragraphs, path):
    path.write_text("\n\n".join([spec.topic.upper()] + paragraphs) + "\n", encoding="utf-8")


def _pdf(spec, paragraphs, path):
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 8, _latin1(spec.topic), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font("Helvetica", size=11)
    for p in paragraphs:
        pdf.ln(2)
        pdf.multi_cell(0, 6, _latin1(p), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.output(str(path))


def _docx(spec, paragraphs, path):
    from docx import Document

    doc = Document()
    doc.add_heading(spec.topic, 1)
    for p in paragraphs:
        doc.add_paragraph(p)
    doc.save(path)


def _pptx(spec, paragraphs, path):
    from pptx import Presentation

    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = spec.topic
    for i in range(0, len(paragraphs), 3):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec.topic
        slide.placeholders[1].text_frame.text = "\n".join(paragraphs[i : i + 3])
    prs.save(path)


def _xlsx(spec, paragraphs, path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    for row in spec.table or [(spec.topic,)]:
        ws.append(list(row))
    notes = wb.create_sheet("Notes")
    notes.append([spec.topic])
    for p in paragraphs:
        notes.append([p])
    wb.save(path)


WRITERS = {"md": _md, "txt": _txt, "pdf": _pdf, "docx": _docx, "pptx": _pptx, "xlsx": _xlsx}
