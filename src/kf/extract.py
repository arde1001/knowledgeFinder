"""Plain-text extraction for the corpus formats (md/txt, pdf, docx, xlsx, pptx)."""

from pathlib import Path

SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".xlsx", ".pptx"}


def extract_text(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".md", ".txt"):
        return path.read_text(encoding="utf-8")
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            lines += ["\t".join(c.text for c in row.cells) for row in table.rows]
        return "\n".join(lines)
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[{ws.title}]")
            lines += ["\t".join("" if c is None else str(c) for c in row) for row in ws.iter_rows(values_only=True)]
        return "\n".join(lines)
    if suffix == ".pptx":
        from pptx import Presentation

        lines = []
        for i, slide in enumerate(Presentation(path).slides, 1):
            lines.append(f"[slide {i}]")
            lines += [s.text_frame.text for s in slide.shapes if s.has_text_frame]
        return "\n".join(lines)
    raise ValueError(f"unsupported file type: {path.suffix}")
